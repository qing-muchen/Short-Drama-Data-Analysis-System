"""Safe in-memory parsing for fixed-template CSV and XLSX files."""

import csv
import io
import zipfile
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import Any

from openpyxl import load_workbook

from imports.templates import MAX_IMPORT_FILE_SIZE_BYTES, MAX_IMPORT_ROWS, get_template


class ImportFileError(ValueError):
    """Safe file-level error suitable for an ImportError record."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


@dataclass(frozen=True, slots=True)
class ParsedRow:
    source_row_number: int
    values: dict[str, Any]


@dataclass(frozen=True, slots=True)
class ParsedFile:
    original_filename: str
    file_format: str
    headers: tuple[str, ...]
    rows: tuple[ParsedRow, ...]


def sanitize_filename(filename: str) -> str:
    candidate = filename.strip()
    normalized = candidate.replace("\\", "/")
    basename = PurePosixPath(normalized).name
    if not candidate or basename != normalized or basename in {".", ".."}:
        raise ImportFileError("UNSAFE_FILENAME", "The upload filename is not safe")
    if len(basename) > 255 or "\x00" in basename:
        raise ImportFileError("UNSAFE_FILENAME", "The upload filename is not safe")
    return basename


def parse_import_file(
    *,
    filename: str,
    file_format: str,
    content: bytes,
    import_type: str,
) -> ParsedFile:
    safe_name = sanitize_filename(filename)
    normalized_format = file_format.upper()
    expected_suffix = ".csv" if normalized_format == "CSV" else ".xlsx"
    if normalized_format not in {"CSV", "XLSX"} or not safe_name.lower().endswith(
        expected_suffix
    ):
        raise ImportFileError(
            "UNSUPPORTED_FILE_TYPE", "Only matching .csv and .xlsx files are supported"
        )
    if not content:
        raise ImportFileError("EMPTY_FILE", "The import file is empty")
    if len(content) > MAX_IMPORT_FILE_SIZE_BYTES:
        raise ImportFileError("FILE_TOO_LARGE", "The import file exceeds 20 MiB")

    if normalized_format == "CSV":
        headers, rows = _parse_csv(content)
    else:
        headers, rows = _parse_xlsx(content)

    template = get_template(import_type)
    missing = sorted(template.required_columns - set(headers))
    unexpected = sorted(set(headers) - set(template.columns))
    if missing or unexpected:
        details = []
        if missing:
            details.append(f"missing columns: {', '.join(missing)}")
        if unexpected:
            details.append(f"unexpected columns: {', '.join(unexpected)}")
        raise ImportFileError("TEMPLATE_MISMATCH", "; ".join(details))
    if not rows:
        raise ImportFileError("EMPTY_DATA", "The import file contains no data rows")
    return ParsedFile(safe_name, normalized_format, headers, rows)


def _normalize_headers(values: list[Any]) -> tuple[str, ...]:
    headers = tuple(str(value).strip().lower() if value is not None else "" for value in values)
    if not headers or any(not header for header in headers):
        raise ImportFileError("INVALID_HEADER", "Headers must be non-empty")
    if len(set(headers)) != len(headers):
        raise ImportFileError("DUPLICATE_HEADER", "Duplicate headers are not allowed")
    return headers


def _rows_from_values(headers: tuple[str, ...], source_rows: Any) -> tuple[ParsedRow, ...]:
    rows: list[ParsedRow] = []
    for row_number, raw_values in source_rows:
        values = list(raw_values)
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        if len(values) > len(headers) and any(
            value is not None and str(value).strip() for value in values[len(headers) :]
        ):
            raise ImportFileError("ROW_WIDTH_MISMATCH", f"Row {row_number} has extra values")
        values.extend([None] * (len(headers) - len(values)))
        rows.append(ParsedRow(row_number, dict(zip(headers, values[: len(headers)]))))
        if len(rows) > MAX_IMPORT_ROWS:
            raise ImportFileError(
                "TOO_MANY_ROWS", f"The import file exceeds {MAX_IMPORT_ROWS} data rows"
            )
    return tuple(rows)


def _parse_csv(content: bytes) -> tuple[tuple[str, ...], tuple[ParsedRow, ...]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise ImportFileError("INVALID_ENCODING", "CSV must use UTF-8 encoding") from error
    try:
        reader = csv.reader(io.StringIO(text, newline=""))
        first = next(reader, None)
        if first is None:
            raise ImportFileError("EMPTY_FILE", "The import file is empty")
        headers = _normalize_headers(first)
        rows = _rows_from_values(headers, enumerate(reader, start=2))
        return headers, rows
    except csv.Error as error:
        raise ImportFileError("CSV_PARSE_ERROR", "CSV content could not be parsed") from error


def _parse_xlsx(content: bytes) -> tuple[tuple[str, ...], tuple[ParsedRow, ...]]:
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise ImportFileError("XLSX_PARSE_ERROR", "The file is not a valid XLSX workbook")
    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        first = next(iterator, None)
        if first is None:
            raise ImportFileError("EMPTY_FILE", "The import file is empty")
        headers = _normalize_headers(list(first))
        rows = _rows_from_values(headers, enumerate(iterator, start=2))
        return headers, rows
    except ImportFileError:
        raise
    except Exception as error:
        raise ImportFileError("XLSX_PARSE_ERROR", "XLSX content could not be parsed") from error
    finally:
        if "workbook" in locals():
            workbook.close()
