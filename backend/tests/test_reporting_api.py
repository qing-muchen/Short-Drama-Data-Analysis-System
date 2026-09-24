"""HTTP boundary tests for the Phase 10 reporting API."""

from __future__ import annotations

import ast
import json
from contextlib import contextmanager
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from auth.exceptions import AuthorizationException
from auth.dependencies import require_user
from main import app
from routers.dependencies import (
    ReportingApiService,
    get_reporting_service,
)
from services.reporting.contracts import (
    AnalysisSnapshot,
    ExportFormat,
    ReportAvailability,
    ReportAvailabilityReason,
    ReportAvailabilityStatus,
    ReportColumnType,
    ReportMetadata,
    ReportResponse,
    ReportRow,
    ReportType,
)
from services.reporting.policies import ExportLimitPolicy
from services.reporting.registry import create_default_registry
from services.reporting.serializers.csv_serializer import CsvReportSerializer
from services.reporting.serializers.xlsx_serializer import XlsxReportSerializer
from services.reporting.service import ReportExportPreparation


AS_OF = datetime(2026, 9, 24, 2, 0, tzinfo=UTC)
CATALOG_PATH = "/api/v1/reports/catalog"
PREVIEW_PATH = "/api/v1/reports/preview"
EXPORT_PATH = "/api/v1/reports/export"


def _value(column_type: ReportColumnType, name: str, nullable: bool):
    if nullable:
        return None
    if name == "drama_name":
        return "=SUM(A1:A2)"
    return {
        ReportColumnType.STRING: "value",
        ReportColumnType.INTEGER: 1,
        ReportColumnType.DECIMAL: Decimal("1.25"),
        ReportColumnType.BOOLEAN: True,
        ReportColumnType.DATE: date(2026, 9, 24),
        ReportColumnType.DATETIME: AS_OF,
    }[column_type]


def _response(
    report_type: ReportType,
    *,
    include_rows: bool = True,
    availability: ReportAvailability | None = None,
    total_rows: int = 1,
) -> ReportResponse:
    schema = create_default_registry().get_schema(report_type)
    row = ReportRow(
        tuple(
            (
                column.name,
                _value(column.data_type, column.name, column.nullable),
            )
            for column in schema.columns
        )
    )
    snapshot = AnalysisSnapshot(
        requested_at=AS_OF,
        analysis_as_of=AS_OF,
        semantics=schema.snapshot_semantics,
    )
    return ReportResponse(
        metadata=ReportMetadata(
            report_type=report_type,
            title=schema.name,
            description=schema.description,
            schema_version=schema.schema_version,
            snapshot=snapshot,
            total_rows=total_rows,
        ),
        columns=schema.columns,
        rows=(row,) if include_rows else (),
        availability=availability
        or ReportAvailability(ReportAvailabilityStatus.AVAILABLE),
    )


class FakeCoreReportingService:
    def __init__(self) -> None:
        self.registry = create_default_registry()
        self.preview_calls: list[dict[str, object]] = []
        self.export_calls: list[dict[str, object]] = []
        self.availability = ReportAvailability(ReportAvailabilityStatus.AVAILABLE)
        self.policy = ExportLimitPolicy()
        self.total_rows = 1

    def get_catalog(self):
        return self.registry.list_schemas()

    def preview_report(self, **kwargs):
        self.preview_calls.append(kwargs)
        return _response(
            kwargs["report_type"],
            availability=self.availability,
        )

    def prepare_export(self, **kwargs):
        self.export_calls.append(kwargs)
        report_type = kwargs["report_type"]
        export_format = kwargs["export_format"]
        decision = self.policy.evaluate_rows(
            total_rows=self.total_rows,
            export_format=export_format,
            has_effective_filters=bool(kwargs["filters"]),
            confirmed_row_count=kwargs["confirmed_row_count"],
        )
        schema = self.registry.get_schema(report_type)
        descriptor = next(
            item
            for item in schema.export_descriptors
            if item.format is export_format
        )
        return ReportExportPreparation(
            report=_response(
                report_type,
                include_rows=decision.allowed,
                availability=self.availability,
                total_rows=self.total_rows,
            ),
            schema=schema,
            descriptor=descriptor,
            decision=decision,
        )


def _api_service(core: FakeCoreReportingService) -> ReportingApiService:
    return ReportingApiService(
        core,  # type: ignore[arg-type]
        {
            ExportFormat.CSV: CsvReportSerializer(),
            ExportFormat.XLSX: XlsxReportSerializer(),
        },
        ExportLimitPolicy(),
    )


@contextmanager
def api_client(*, role: str | None = "USER"):
    core = FakeCoreReportingService()
    service = _api_service(core)
    app.dependency_overrides[get_reporting_service] = lambda: service
    if role == "FORBIDDEN":
        def denied():
            raise AuthorizationException()

        app.dependency_overrides[require_user] = denied
    elif role is not None:
        app.dependency_overrides[require_user] = lambda: SimpleNamespace(
            user_id=1,
            role_code=role,
        )
    try:
        yield TestClient(app), core
    finally:
        app.dependency_overrides.pop(get_reporting_service, None)
        app.dependency_overrides.pop(require_user, None)


def _authorized_request(client: TestClient, path: str):
    if path == CATALOG_PATH:
        return client.get(path)
    if path == PREVIEW_PATH:
        return client.get(path, params={"report_type": "DRAMA_CATALOG"})
    return client.get(
        path,
        params={
            "report_type": "DRAMA_CATALOG",
            "export_format": "CSV",
            "filters": json.dumps({"drama_id": 1}),
        },
    )


def test_three_routes_are_registered_as_get_only() -> None:
    paths = app.openapi()["paths"]
    for path in (CATALOG_PATH, PREVIEW_PATH, EXPORT_PATH):
        assert set(paths[path]) == {"get"}


@pytest.mark.parametrize("path", (CATALOG_PATH, PREVIEW_PATH, EXPORT_PATH))
def test_anonymous_requests_return_401(path: str) -> None:
    with api_client(role=None) as (client, _core):
        response = _authorized_request(client, path)

    assert response.status_code == 401
    assert response.json()["code"] == "AUTHENTICATION_FAILED"


@pytest.mark.parametrize("role", ("USER", "ADMINISTRATOR"))
def test_user_and_administrator_can_access_all_routes(role: str) -> None:
    with api_client(role=role) as (client, _core):
        responses = [
            _authorized_request(client, path)
            for path in (CATALOG_PATH, PREVIEW_PATH, EXPORT_PATH)
        ]

    assert [response.status_code for response in responses] == [200, 200, 200]


def test_forbidden_identity_returns_403() -> None:
    with api_client(role="FORBIDDEN") as (client, _core):
        response = client.get(CATALOG_PATH)

    assert response.status_code == 403
    assert response.json()["code"] == "FORBIDDEN"


def test_catalog_returns_schema_versions_columns_filters_and_formats() -> None:
    with api_client() as (client, _core):
        response = client.get(CATALOG_PATH)

    assert response.status_code == 200
    catalog = response.json()["data"]
    assert len(catalog) == 6
    drama = catalog[0]
    assert drama["report_type"] == "DRAMA_CATALOG"
    assert drama["schema_version"] == "DRAMA_CATALOG Schema V1.0"
    assert drama["columns"][0]["name"] == "drama_id"
    assert drama["supported_formats"] == ["CSV", "XLSX"]
    assert {item["key"] for item in drama["filters"]} == {
        "drama_id",
        "platform_id",
        "status",
        "tag_id",
    }


def test_preview_validates_and_forwards_filters_cutoff_and_page_size() -> None:
    filters = {"drama_id": 7, "status": "COMPLETED"}
    with api_client() as (client, core):
        response = client.get(
            PREVIEW_PATH,
            params={
                "report_type": "DRAMA_CATALOG",
                "filters": json.dumps(filters),
                "analysis_as_of": AS_OF.isoformat(),
                "page_size": 37,
            },
        )

    assert response.status_code == 200
    assert core.preview_calls == [
        {
            "report_type": ReportType.DRAMA_CATALOG,
            "filters": filters,
            "analysis_as_of": AS_OF,
            "limit": 37,
        }
    ]
    data = response.json()["data"]
    assert data["metadata"]["snapshot"]["analysis_as_of"] == "2026-09-24T02:00:00Z"
    assert data["columns"][0]["name"] == "drama_id"
    assert list(data["rows"][0]) == [
        column["name"] for column in data["columns"]
    ]


@pytest.mark.parametrize("page_size", (0, 101))
def test_preview_page_size_is_limited(page_size: int) -> None:
    with api_client() as (client, _core):
        response = client.get(
            PREVIEW_PATH,
            params={
                "report_type": "DRAMA_CATALOG",
                "page_size": page_size,
            },
        )
    assert response.status_code == 422


def test_invalid_filter_json_returns_safe_400() -> None:
    with api_client() as (client, _core):
        response = client.get(
            PREVIEW_PATH,
            params={"report_type": "DRAMA_CATALOG", "filters": "not-json"},
        )

    assert response.status_code == 400
    assert response.json()["code"] == "INVALID_REPORT_FILTERS"
    assert "traceback" not in response.text.lower()


@pytest.mark.parametrize(
    "status",
    tuple(ReportAvailabilityStatus),
)
def test_preview_preserves_all_availability_states(
    status: ReportAvailabilityStatus,
) -> None:
    reasons = ()
    if status is not ReportAvailabilityStatus.AVAILABLE:
        reasons = (
            ReportAvailabilityReason.NO_MATCHING_DATA
            if status is ReportAvailabilityStatus.EMPTY
            else ReportAvailabilityReason.MISSING_SOURCE,
        )
    with api_client() as (client, core):
        core.availability = ReportAvailability(status, reasons)
        response = client.get(
            PREVIEW_PATH,
            params={"report_type": "DRAMA_CATALOG"},
        )

    assert response.status_code == 200
    assert response.json()["data"]["availability"]["status"] == status.value


def test_csv_export_streams_safe_mime_headers_bom_and_formula_escape() -> None:
    with api_client() as (client, _core):
        response = client.get(
            EXPORT_PATH,
            params={
                "report_type": "DRAMA_CATALOG",
                "export_format": "CSV",
                "filters": json.dumps({"drama_id": 1}),
            },
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert response.headers["content-disposition"].startswith("attachment;")
    assert response.headers["x-report-schema-version"]
    assert response.headers["x-analysis-as-of"] == AS_OF.isoformat()
    assert response.content.startswith(b"\xef\xbb\xbf")
    assert b"'=SUM(A1:A2)" in response.content


def test_xlsx_export_streams_correct_mime_and_workbook() -> None:
    with api_client() as (client, _core):
        response = client.get(
            EXPORT_PATH,
            params={
                "report_type": "DRAMA_CATALOG",
                "export_format": "XLSX",
                "filters": json.dumps({"drama_id": 1}),
            },
        )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content), read_only=True)
    try:
        assert workbook.sheetnames == ["Drama Catalog"]
        assert workbook["Drama Catalog"]["C2"].value == "'=SUM(A1:A2)"
    finally:
        workbook.close()


def test_export_confirmation_is_required_and_then_forwarded() -> None:
    with api_client() as (client, core):
        initial = client.get(
            EXPORT_PATH,
            params={
                "report_type": "DRAMA_CATALOG",
                "export_format": "CSV",
            },
        )
        confirmed = client.get(
            EXPORT_PATH,
            params={
                "report_type": "DRAMA_CATALOG",
                "export_format": "CSV",
                "confirmation": 1,
            },
        )

    assert initial.status_code == 409
    assert initial.json()["code"] == "EXPORT_CONFIRMATION_REQUIRED"
    assert initial.json()["details"]["confirmed_row_count"] == 1
    assert confirmed.status_code == 200
    assert core.export_calls[-1]["confirmed_row_count"] == 1


def test_export_limit_is_returned_as_safe_structured_error() -> None:
    with api_client() as (client, core):
        core.total_rows = 10_001
        response = client.get(
            EXPORT_PATH,
            params={
                "report_type": "DRAMA_CATALOG",
                "export_format": "CSV",
                "filters": json.dumps({"drama_id": 1}),
            },
        )

    assert response.status_code == 400
    assert response.json()["code"] == "EXPORT_LIMIT_EXCEEDED"
    assert response.json()["details"] == {
        "total_rows": 10_001,
        "max_rows": 10_000,
    }


def test_unavailable_export_returns_structured_error_instead_of_stream() -> None:
    availability = ReportAvailability(
        ReportAvailabilityStatus.NOT_AVAILABLE,
        (ReportAvailabilityReason.MISSING_SOURCE,),
    )
    with api_client() as (client, core):
        core.availability = availability
        response = client.get(
            EXPORT_PATH,
            params={
                "report_type": "DRAMA_CATALOG",
                "export_format": "CSV",
                "filters": json.dumps({"drama_id": 1}),
            },
        )

    assert response.status_code == 400
    assert response.json()["code"] == "REPORT_NOT_AVAILABLE"
    assert response.json()["details"] == {"availability": "NOT_AVAILABLE"}


def test_dependency_composes_request_scoped_reporting_application() -> None:
    result = get_reporting_service(object())  # type: ignore[arg-type]
    assert isinstance(result, ReportingApiService)


def test_router_has_no_forbidden_data_or_export_dependencies() -> None:
    path = Path(__file__).parents[1] / "routers" / "reporting.py"
    tree = ast.parse(path.read_text(encoding="utf-8"))
    forbidden_roots = {"database", "models", "repositories", "sqlalchemy"}
    forbidden_parts = {"providers", "registry", "serializers"}
    forbidden_calls = {"commit", "rollback", "flush", "query", "execute"}

    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            for alias in node.names:
                assert alias.name.split(".")[0] not in forbidden_roots
                assert not set(alias.name.split(".")).intersection(forbidden_parts)
        if isinstance(node, ast.ImportFrom) and node.module:
            parts = set(node.module.split("."))
            assert node.module.split(".")[0] not in forbidden_roots
            assert not parts.intersection(forbidden_parts)
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
            assert node.func.attr not in forbidden_calls
