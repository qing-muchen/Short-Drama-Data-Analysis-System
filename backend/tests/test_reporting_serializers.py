"""Safety and architecture tests for report transport serializers."""

from __future__ import annotations

import ast
import csv
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from services.reporting.serializers.csv_serializer import (
    CSV_MIME_TYPE,
    CsvReportSerializer,
)
from services.reporting.serializers.xlsx_serializer import (
    XLSX_MIME_TYPE,
    XlsxReportSerializer,
)


@dataclass(frozen=True)
class Column:
    key: str
    title: str
    column_type: str
    nullable: bool = False


@dataclass(frozen=True)
class Schema:
    columns: tuple[Column, ...]
    sheet_name: str = "Data"


SCHEMA = Schema(
    (
        Column("name", "Name", "STRING"),
        Column("count", "Count", "INTEGER"),
        Column("ratio", "Ratio", "DECIMAL"),
        Column("active", "Active", "BOOLEAN"),
        Column("business_date", "Date", "DATE"),
        Column("analysis_as_of", "As Of", "DATETIME"),
        Column("optional", "Optional", "STRING", nullable=True),
    )
)


def _row(name: str = "短剧, \"甲\"\n第二行") -> dict[str, object]:
    return {
        "name": name,
        "count": 0,
        "ratio": Decimal("1.25"),
        "active": False,
        "business_date": date(2026, 9, 23),
        "analysis_as_of": datetime(2026, 9, 23, 8, tzinfo=UTC),
        "optional": None,
    }


def test_csv_has_bom_crlf_utf8_quoting_and_null_semantics() -> None:
    with CsvReportSerializer().serialize([_row()], SCHEMA) as result:
        content = b"".join(result.iter_bytes())
        assert result.mime_type == CSV_MIME_TYPE
        assert result.file_extension == ".csv"
        assert result.byte_length == len(content)

    assert content.startswith(b"\xef\xbb\xbf")
    assert content.count(b"\r\n") == 2
    decoded = content.decode("utf-8-sig")
    parsed = list(csv.reader(StringIO(decoded, newline="")))
    assert parsed[0] == [column.title for column in SCHEMA.columns]
    assert parsed[1][0] == '短剧, "甲"\n第二行'
    assert parsed[1][1:4] == ["0", "1.25", "false"]
    assert parsed[1][-1] == ""


@pytest.mark.parametrize("prefix", ["=", "+", "-", "*", "@"])
def test_csv_blocks_formula_injection_including_leading_whitespace(prefix: str) -> None:
    with CsvReportSerializer().serialize(
        [_row(f"  {prefix}SUM(A1:A2)")], SCHEMA
    ) as result:
        parsed = list(
            csv.reader(StringIO(b"".join(result.iter_bytes()).decode("utf-8-sig")))
        )
    assert parsed[1][0] == f"'  {prefix}SUM(A1:A2)"


def test_csv_rejects_none_for_non_nullable_and_schema_mismatch() -> None:
    invalid = _row()
    invalid["count"] = None
    with pytest.raises(ValueError, match="non-nullable"):
        CsvReportSerializer().serialize([invalid], SCHEMA)
    with pytest.raises(ValueError, match="row/schema mismatch"):
        CsvReportSerializer().serialize([{"name": "only"}], SCHEMA)


def test_csv_rolls_large_output_out_of_memory() -> None:
    serializer = CsvReportSerializer(spool_threshold_bytes=64)
    rows = (_row(f"row-{index}-" + "x" * 40) for index in range(20))
    with serializer.serialize(rows, SCHEMA) as result:
        assert result.rolled_to_disk is True
        assert result.byte_length > 64


def test_csv_is_deterministic_for_the_same_schema_and_rows() -> None:
    serializer = CsvReportSerializer()
    with serializer.serialize([_row()], SCHEMA) as first:
        first_bytes = b"".join(first.iter_bytes())
    with serializer.serialize([_row()], SCHEMA) as second:
        second_bytes = b"".join(second.iter_bytes())
    assert first_bytes == second_bytes


def test_xlsx_generation_mime_values_and_write_only_safety() -> None:
    with XlsxReportSerializer().serialize([_row("=1+1")], SCHEMA) as result:
        payload = b"".join(result.iter_bytes())
        assert result.mime_type == XLSX_MIME_TYPE
        assert result.file_extension == ".xlsx"

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["Data"]
        rows = list(workbook["Data"].iter_rows(min_col=1, max_col=7, values_only=False))
        assert [cell.value for cell in rows[0]] == [
            column.title for column in SCHEMA.columns
        ]
        assert rows[1][0].value == "'=1+1"
        assert rows[1][0].data_type == "s"
        assert rows[1][1].value == 0
        assert rows[1][-1].value is None
        assert not getattr(workbook, "_external_links", [])
    finally:
        workbook.close()


def test_xlsx_uses_spool_and_accepts_streaming_iterable() -> None:
    serializer = XlsxReportSerializer(spool_threshold_bytes=64)
    rows = (_row(f"row-{index}") for index in range(200))
    with serializer.serialize(rows, SCHEMA) as result:
        assert result.rolled_to_disk is True
        assert result.byte_length > 64


def test_xlsx_rejects_unsafe_sheet_name_and_naive_datetime() -> None:
    with pytest.raises(ValueError, match="sheet name"):
        XlsxReportSerializer().serialize([_row()], Schema(SCHEMA.columns, "Bad/Name"))
    invalid = _row()
    invalid["analysis_as_of"] = datetime(2026, 9, 23)
    with pytest.raises(ValueError, match="timezone-aware"):
        XlsxReportSerializer().serialize([invalid], SCHEMA)


def test_serializers_close_request_scoped_streams() -> None:
    result = CsvReportSerializer().serialize([_row()], SCHEMA)
    result.close()
    assert result.stream.closed is True


def test_serializer_modules_have_no_data_layer_or_service_dependencies() -> None:
    root = Path(__file__).parents[1] / "services" / "reporting" / "serializers"
    forbidden_import_roots = {"database", "models", "repositories", "sqlalchemy"}
    forbidden_calls = {"commit", "rollback", "flush", "query", "execute"}

    for path in root.glob("*.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                assert not any(
                    alias.name.split(".")[0] in forbidden_import_roots
                    for alias in node.names
                )
            if isinstance(node, ast.ImportFrom) and node.module:
                assert node.module.split(".")[0] not in forbidden_import_roots
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
                assert node.func.attr not in forbidden_calls
