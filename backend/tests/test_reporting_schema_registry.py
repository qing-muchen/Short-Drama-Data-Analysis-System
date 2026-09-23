"""Schema registry, policy, and serializer compatibility tests."""

from __future__ import annotations

import ast
import csv
from datetime import UTC, date, datetime
from io import BytesIO, StringIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from services.reporting.contracts import (
    ExportFormat,
    ReportRow,
    ReportType,
    SnapshotSemantics,
)
from services.reporting.policies import ReportFilterPolicy
from services.reporting.registry import (
    ReportSchemaRegistry,
    UnknownReportTypeError,
    create_default_registry,
)
from services.reporting.serializers.csv_serializer import CsvReportSerializer
from services.reporting.serializers.xlsx_serializer import XlsxReportSerializer


def test_default_registry_contains_exactly_six_explicit_reports() -> None:
    registry = create_default_registry()
    assert len(registry) == 6
    assert tuple(schema.report_type for schema in registry.list_schemas()) == tuple(
        ReportType
    )
    assert all(schema.identifier == schema.report_type.value for schema in registry.list_schemas())


def test_every_schema_has_stable_version_columns_formats_and_descriptors() -> None:
    registry = create_default_registry()
    for schema in registry.list_schemas():
        assert schema.schema_version == f"{schema.report_type.value} Schema V1.0"
        assert schema.columns
        assert len({column.name for column in schema.columns}) == len(schema.columns)
        assert schema.supported_formats == (ExportFormat.CSV, ExportFormat.XLSX)
        assert tuple(item.format for item in schema.export_descriptors) == (
            ExportFormat.CSV,
            ExportFormat.XLSX,
        )
        assert schema.name
        assert schema.description


def test_snapshot_semantics_match_frozen_capability_matrix() -> None:
    registry = create_default_registry()
    as_of_types = {
        ReportType.SCORE_RESULTS,
        ReportType.DATA_QUALITY_RESULTS,
        ReportType.PREDICTION_RESULTS,
    }
    for report_type in ReportType:
        expected = (
            SnapshotSemantics.AS_OF_RECONSTRUCTED
            if report_type in as_of_types
            else SnapshotSemantics.CURRENT_STATE_AT_REQUEST
        )
        assert registry.get_schema(report_type).snapshot_semantics is expected


def test_column_order_is_explicit_and_stable() -> None:
    registry = create_default_registry()
    score = registry.get_schema(ReportType.SCORE_RESULTS)
    assert tuple(column.name for column in score.columns) == (
        "score_id",
        "drama_id",
        "score_date",
        "score_scope_type",
        "score_scope_id",
        "analysis_as_of",
        "heat_score",
        "growth_score",
        "engagement_score",
        "lifecycle_score",
        "total_score",
        "level",
        "algorithm_version",
    )


def test_filter_definitions_drive_the_existing_filter_policy() -> None:
    registry = create_default_registry()
    result = ReportFilterPolicy(registry).validate(
        ReportType.ANOMALY_RECORDS,
        {
            "drama_id": 9,
            "status": "OPEN",
            "start_time": "2026-09-01T00:00:00Z",
        },
    )
    assert result.as_mapping()["status"] == "OPEN"
    assert result.as_mapping()["start_time"] == datetime(
        2026, 9, 1, tzinfo=UTC
    )
    with pytest.raises(ValueError):
        ReportFilterPolicy(registry).validate(
            ReportType.DATA_QUALITY_RESULTS,
            {"drama_id": 1},
        )


def test_registry_rejects_duplicate_and_unknown_types() -> None:
    schema = create_default_registry().get_schema(ReportType.DRAMA_CATALOG)
    with pytest.raises(ValueError, match="duplicate"):
        ReportSchemaRegistry((schema, schema))
    with pytest.raises(UnknownReportTypeError):
        create_default_registry().get_schema("UNKNOWN")  # type: ignore[arg-type]


def test_real_schema_and_row_are_csv_serializer_compatible() -> None:
    schema = create_default_registry().get_schema(ReportType.DRAMA_CATALOG)
    row = ReportRow(
        (
            ("drama_id", 1),
            ("drama_code", "D-001"),
            ("drama_name", "=Unsafe Name"),
            ("status", "COMPLETED"),
            ("premiere_date", date(2026, 9, 1)),
            ("episode_count", 10),
            ("episode_duration_seconds", 90),
            ("platforms", "Platform A"),
            ("tags", None),
        )
    )
    with CsvReportSerializer().serialize((row,), schema) as result:
        parsed = list(
            csv.reader(StringIO(b"".join(result.iter_bytes()).decode("utf-8-sig")))
        )
    assert parsed[0] == [column.display_name for column in schema.columns]
    assert parsed[1][2] == "'=Unsafe Name"
    assert parsed[1][-1] == ""


def test_real_schema_and_row_are_xlsx_serializer_compatible() -> None:
    schema = create_default_registry().get_schema(ReportType.DRAMA_CATALOG)
    row = ReportRow(
        (
            ("drama_id", 1),
            ("drama_code", "D-001"),
            ("drama_name", "Drama"),
            ("status", "COMPLETED"),
            ("premiere_date", None),
            ("episode_count", None),
            ("episode_duration_seconds", None),
            ("platforms", None),
            ("tags", None),
        )
    )
    with XlsxReportSerializer().serialize((row,), schema) as result:
        payload = b"".join(result.iter_bytes())
    workbook = load_workbook(BytesIO(payload), read_only=True)
    try:
        assert workbook.sheetnames == [schema.sheet_name]
        values = list(
            workbook[schema.sheet_name].iter_rows(
                min_col=1,
                max_col=len(schema.columns),
                values_only=True,
            )
        )
        assert values[1][0:4] == (1, "D-001", "Drama", "COMPLETED")
        assert values[1][4:] == (None, None, None, None, None)
    finally:
        workbook.close()


def test_registry_is_deterministic_and_has_no_mutable_global_instance() -> None:
    first = create_default_registry()
    second = create_default_registry()
    assert first is not second
    assert first.list_schemas() == second.list_schemas()
    assert not hasattr(first, "register")


def test_registry_modules_have_no_data_or_application_layer_dependencies() -> None:
    root = Path(__file__).parents[1] / "services" / "reporting" / "registry"
    forbidden_roots = {
        "database",
        "fastapi",
        "models",
        "repositories",
        "sqlalchemy",
    }
    forbidden_calls = {"commit", "rollback", "flush", "query", "execute"}
    for path in root.glob("*.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        for node in ast.walk(tree):
            if isinstance(node, ast.Import):
                assert not any(
                    alias.name.split(".")[0] in forbidden_roots
                    for alias in node.names
                )
            if isinstance(node, ast.ImportFrom) and node.module:
                assert node.module.split(".")[0] not in forbidden_roots
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute):
                assert node.func.attr not in forbidden_calls
